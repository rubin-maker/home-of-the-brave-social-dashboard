#!/usr/bin/env python3
"""Normalize the four Home of the Brave workbooks into one dashboard data layer."""

import csv
import json
import os
from collections import defaultdict
from datetime import date, datetime, timedelta

try:
    from openpyxl import load_workbook
    from openpyxl.utils.datetime import from_excel
except ImportError:
    print("openpyxl is required — run: python3 -m pip install -r requirements.txt")
    raise SystemExit(1)


ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "sources")
BRAND = "Home of the Brave"
WEEKLY_CONTEXT_WEEKS = 8
CATEGORY_TREND_PLATFORMS = ("Instagram", "YouTube", "TikTok")

SOURCES = {
    "YouTube": os.path.join(SRC, "youtube.xlsx"),
    "Instagram": os.path.join(SRC, "instagram.xlsx"),
    "TikTok": os.path.join(SRC, "tiktok.xlsx"),
    "X": os.path.join(SRC, "x.xlsx"),
}
YOUTUBE_LONG_SOURCE = os.path.join(SRC, "youtube_long.csv")
YOUTUBE_SHORTS_LIKES_SOURCE = os.path.join(SRC, "youtube_shorts_likes.csv")
# The two update files were exported on Sep 18. The daily Shorts activity file
# ends Sep 17, which is consistent with an export produced the following day.
YOUTUBE_UPDATE_THROUGH = "2026-09-18"

METRIC_DEFINITIONS = {
    "YouTube": {
        "views": "Views",
        "engagements": "Likes + shares + comments",
        "audience": "Impressions",
        "metric_label": "views",
        "subscribers": "Source-reported Subscribers attributed to exported videos; not a current channel subscriber count or net subscriber growth.",
    },
    "Instagram": {
        "views": "Views",
        "engagements": "Likes + comments + shares + saves",
        "audience": "Reach",
        "metric_label": "views",
    },
    "TikTok": {
        "views": "Views",
        "engagements": "Likes + comments + shares + bookmarks",
        "audience": "Not available",
        "metric_label": "views",
    },
    "X": {
        "views": "Impressions (reported as Views in the source workbook)",
        "engagements": "Source-reported total engagements",
        "audience": "Not available",
        "metric_label": "impressions",
    },
}


def clean_text(value):
    text = "\n".join(line.rstrip() for line in str(value or "").splitlines()).strip()
    replacements = {
        "‚Äú": "“", "‚Äù": "”", "‚Äô": "’", "‚Ä¶": "…",
        "â€œ": "“", "â€": "”", "â€™": "’", "â€¦": "…",
    }
    for bad, good in replacements.items():
        text = text.replace(bad, good)
    return text


def number(value):
    if value in (None, "", "-", "--", "N/A"):
        return 0.0
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def iso_date(value, epoch):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)):
        return from_excel(value, epoch).date().isoformat()
    text = clean_text(value)
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%Y %H:%M", "%b %d, %Y"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            pass
    raise ValueError(f"Unrecognized date: {value!r}")


def monday_of(iso):
    current = date.fromisoformat(iso)
    return current - timedelta(days=current.weekday())


def read_rows(workbook_path, sheet_name):
    wb = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        ws = wb[sheet_name]
        iterator = ws.iter_rows(values_only=True)
        headers = [clean_text(v) for v in next(iterator)]
        for values in iterator:
            if not any(v is not None for v in values):
                continue
            yield dict(zip(headers, values)), wb.epoch
    finally:
        wb.close()


def read_csv_rows(path):
    with open(path, encoding="utf-8-sig", newline="") as handle:
        yield from csv.DictReader(handle)


def workbook_definition_date(path, labels):
    """Return the first matching Definitions date without trusting overview copy."""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb["Definitions"]
        for row in ws.iter_rows(values_only=True):
            if clean_text(row[0]) in labels and len(row) > 1 and row[1] is not None:
                return iso_date(row[1], wb.epoch)
    finally:
        wb.close()
    raise ValueError(f"No Definitions date {sorted(labels)} in {path}")


def youtube_posts():
    posts = []

    # The reviewed Sep 10 workbook remains the category authority and the full
    # Shorts catalog. The Sep 18 long-form CSV replaces its long-form metrics.
    long_categories = {}
    for row, _epoch in read_rows(SOURCES["YouTube"], "Long Data"):
        content_id = clean_text(row.get("Content"))
        category = clean_text(row.get("Category"))
        if content_id and category:
            long_categories[content_id] = category

    for row in read_csv_rows(YOUTUBE_LONG_SOURCE):
        content_id = clean_text(row.get("Content"))
        raw_published = clean_text(row.get("Video publish time"))
        if not content_id or content_id.lower() == "total" or not raw_published:
            continue
        published = iso_date(raw_published, None)
        if not published.startswith("2026-"):
            continue
        title = clean_text(row.get("Video title"))
        category = long_categories.get(content_id)
        if not category and "bill kristol" in title.lower() and "tom joscelyn" in title.lower():
            category = "Bill Kristol & Tom Joscelyn"
        category = category or "Other"
        likes = int(number(row.get("Likes")))
        shares = int(number(row.get("Shares")))
        comments = int(number(row.get("Comments added")))
        posts.append({
            "id": content_id,
            "platform": "YouTube",
            "date": published,
            "title": title,
            "url": f"https://www.youtube.com/watch?v={content_id}",
            "views": int(number(row.get("Views"))),
            "reach": int(number(row.get("Thumbnail impressions"))),
            "engagements": likes + shares + comments,
            "likes": likes,
            "comments": comments,
            "shares": shares,
            "saves": 0,
            "followers": int(number(row.get("Subscribers"))),
            "type": "Long-Form",
            "category": category,
            "duration_sec": int(number(row.get("Duration"))),
            "watch_hours": round(number(row.get("Watch time (hours)")), 4),
        })

    # The supplied Shorts CSV is a daily Likes feed for only five existing
    # videos. Sum every daily adjustment (including negatives), then patch only
    # Likes while retaining the reviewed catalog and all other Sep 10 metrics.
    shorts_likes = defaultdict(int)
    for row in read_csv_rows(YOUTUBE_SHORTS_LIKES_SOURCE):
        content_id = clean_text(row.get("Content"))
        if content_id:
            shorts_likes[content_id] += int(number(row.get("Likes")))

    shorts_ids = set()
    for row, epoch in read_rows(SOURCES["YouTube"], "Shorts Data"):
        if clean_text(row.get("Analysis Status")) != "Included":
            continue
        content_id = clean_text(row.get("Content"))
        shorts_ids.add(content_id)
        published = iso_date(row.get("Video publish time"), epoch)
        likes = shorts_likes.get(content_id, int(number(row.get("Likes"))))
        shares = int(number(row.get("Shares")))
        comments = int(number(row.get("Comments added")))
        posts.append({
            "id": content_id,
            "platform": "YouTube",
            "date": published,
            "title": clean_text(row.get("Video title")),
            "url": f"https://www.youtube.com/watch?v={content_id}" if content_id else "",
            "views": int(number(row.get("Views"))),
            "reach": int(number(row.get("Thumbnail impressions"))),
            "engagements": likes + shares + comments,
            "likes": likes,
            "comments": comments,
            "shares": shares,
            "saves": 0,
            "followers": int(number(row.get("Subscribers"))),
            "type": "Shorts",
            "category": clean_text(row.get("Category")) or "Other",
            "duration_sec": int(number(row.get("Duration"))),
            "watch_hours": round(number(row.get("Watch time (hours)")), 4),
        })

    unknown_short_ids = set(shorts_likes) - shorts_ids
    if unknown_short_ids:
        raise ValueError(f"Shorts Likes feed contains unknown IDs: {sorted(unknown_short_ids)}")
    return posts


def instagram_posts():
    posts = []
    for row, epoch in read_rows(SOURCES["Instagram"], "IG Data"):
        likes = int(number(row.get("Likes")))
        comments = int(number(row.get("Comments")))
        shares = int(number(row.get("Shares")))
        saves = int(number(row.get("Saves")))
        posts.append({
            "id": clean_text(row.get("Post ID")),
            "platform": "Instagram",
            "date": iso_date(row.get("Publish time"), epoch),
            "title": clean_text(row.get("Description")),
            "url": clean_text(row.get("Permalink")),
            "views": int(number(row.get("Views"))),
            "reach": int(number(row.get("Reach"))),
            "engagements": likes + comments + shares + saves,
            "likes": likes,
            "comments": comments,
            "shares": shares,
            "saves": saves,
            "followers": int(number(row.get("Follows"))),
            "type": clean_text(row.get("Post type")),
            "category": clean_text(row.get("Category")) or "Other",
            "duration_sec": int(number(row.get("Duration (sec)"))),
            "watch_hours": 0,
        })
    return posts


def tiktok_posts():
    posts = []
    for row, epoch in read_rows(SOURCES["TikTok"], "TT Data"):
        likes = int(number(row.get("Likes")))
        comments = int(number(row.get("Comments")))
        shares = int(number(row.get("Shares")))
        saves = int(number(row.get("Bookmarks")))
        posts.append({
            "id": clean_text(row.get("Post ID")),
            "platform": "TikTok",
            "date": iso_date(row.get("Publish time"), epoch),
            "title": clean_text(row.get("Post text")),
            "url": clean_text(row.get("URL")),
            "views": int(number(row.get("Views"))),
            "reach": 0,
            # Eight refreshed rows disagree with Total Engagements. Use the
            # documented component definition consistently for every post.
            "engagements": likes + comments + shares + saves,
            "likes": likes,
            "comments": comments,
            "shares": shares,
            "saves": saves,
            "followers": 0,
            "type": "Video",
            "category": clean_text(row.get("Category")) or "Other",
            "duration_sec": int(number(row.get("Duration (sec)"))),
            "watch_hours": 0,
        })
    return posts


def x_posts():
    posts = []
    for row, epoch in read_rows(SOURCES["X"], "X Data"):
        likes = int(number(row.get("Likes")))
        comments = int(number(row.get("Replies")))
        shares = int(number(row.get("Reposts")))
        saves = int(number(row.get("Bookmarks")))
        posts.append({
            "id": clean_text(row.get("Post ID")),
            "platform": "X",
            "date": iso_date(row.get("Date"), epoch),
            "title": clean_text(row.get("Post text")),
            "url": clean_text(row.get("URL")),
            "views": int(number(row.get("Views"))),
            "reach": 0,
            "engagements": int(number(row.get("Engagements"))),
            "likes": likes,
            "comments": comments,
            "shares": shares,
            "saves": saves,
            "followers": 0,
            "type": clean_text(row.get("Source")),
            "category": clean_text(row.get("Category")) or "Other",
            "duration_sec": 0,
            "watch_hours": 0,
        })
    return posts


def aggregate(rows):
    views = sum(p["views"] for p in rows)
    engagements = sum(p["engagements"] for p in rows)
    return {
        "posts": len(rows),
        "views": views,
        "eng": engagements,
        "reach": sum(p["reach"] for p in rows),
        "followers": sum(p["followers"] for p in rows),
        "watch_hours": round(sum(p["watch_hours"] for p in rows), 2),
        "er": round(100 * engagements / views, 2) if views else None,
        "avg_views": round(views / len(rows)) if rows else 0,
    }


posts = youtube_posts() + instagram_posts() + tiktok_posts() + x_posts()
posts.sort(key=lambda p: (p["date"], p["platform"], p["id"]))
platforms = list(SOURCES)

if not posts:
    raise SystemExit("No source rows found.")

period_start = min(p["date"] for p in posts)
period_end = max(p["date"] for p in posts)
platform_latest_posts = {
    platform: max(p["date"] for p in posts if p["platform"] == platform)
    for platform in platforms
}
youtube_complete_through = workbook_definition_date(
    SOURCES["YouTube"], {"Data cutoff", "Data through date"}
)
platform_ends = {
    "YouTube": YOUTUBE_UPDATE_THROUGH,
    "Instagram": workbook_definition_date(SOURCES["Instagram"], {"Data through date"}),
    "TikTok": workbook_definition_date(SOURCES["TikTok"], {"Data through date"}),
    "X": workbook_definition_date(SOURCES["X"], {"Data through date"}),
}
# `platform_ends` is the activity/data-through date for known rows. YouTube's
# complete multi-product catalog is older because the new Shorts file covers
# only Likes for five videos. Keep those concepts separate so missing Shorts
# data is never rendered as a measured zero.
platform_complete_through = {
    **platform_ends,
    "YouTube": youtube_complete_through,
}
for platform in platforms:
    if platform_latest_posts[platform] > platform_ends[platform]:
        raise ValueError(
            f"{platform} latest post {platform_latest_posts[platform]} exceeds "
            f"source coverage {platform_ends[platform]}"
        )
common_coverage_end = min(platform_complete_through.values())
latest_complete_sunday = monday_of(common_coverage_end) - timedelta(days=1)
report_start = latest_complete_sunday - timedelta(days=6)
report_end = latest_complete_sunday

week_starts = []
cursor = monday_of(period_start)
while cursor <= report_start:
    week_starts.append(cursor)
    cursor += timedelta(days=7)
context_starts = week_starts[-WEEKLY_CONTEXT_WEEKS:]
weeks = []
for start in context_starts:
    end = start + timedelta(days=6)
    name = f"Week {start.isocalendar().week}"
    weeks.append([name, start.isoformat(), end.isoformat()])

# Top-post rankings keep the existing comparison context and extend through the
# latest publish week available in any source. Unlike `weeks`, this list is not
# restricted to the last week complete across all four platforms.
top5_weeks = []
top5_cursor = context_starts[0]
top5_last_start = monday_of(max(platform_ends.values()))
while top5_cursor <= top5_last_start:
    top5_end = top5_cursor + timedelta(days=6)
    top5_weeks.append([
        f"Week {top5_cursor.isocalendar().week}",
        top5_cursor.isoformat(),
        top5_end.isoformat(),
    ])
    top5_cursor += timedelta(days=7)

for post in posts:
    post_week_start = monday_of(post["date"])
    post["week"] = f"Week {post_week_start.isocalendar().week}"

# Category charts retain every available 2026 publish week. The first and last
# labels are clipped to the available period. After the partial opening week,
# each aggregation key is the Monday that owns the week.
category_weeks = []
category_cursor = monday_of(period_start)
category_coverage_end = max(platform_ends.values())
category_last_start = monday_of(category_coverage_end)
period_start_date = date.fromisoformat(period_start)
category_coverage_end_date = date.fromisoformat(category_coverage_end)
while category_cursor <= category_last_start:
    category_end = category_cursor + timedelta(days=6)
    display_start = max(category_cursor, period_start_date)
    display_end = min(category_end, category_coverage_end_date)
    category_weeks.append([
        display_start.isoformat(),
        display_start.isoformat(),
        display_end.isoformat(),
    ])
    category_cursor += timedelta(days=7)

totals = {p: aggregate([x for x in posts if x["platform"] == p]) for p in platforms}
youtube_rows = [post for post in posts if post["platform"] == "YouTube"]
youtube_subscribers = {
    "gained": totals["YouTube"]["followers"],
    "long_form": sum(post["followers"] for post in youtube_rows if post["type"] == "Long-Form"),
    "shorts": sum(post["followers"] for post in youtube_rows if post["type"] == "Shorts"),
    "video_count": len(youtube_rows),
    "as_of": YOUTUBE_UPDATE_THROUGH,
    "as_of_label": "Long-form through Sep 18; full Shorts metrics through Sep 10",
    "complete_through": youtube_complete_through,
    "channel_total": None,
    "definition": METRIC_DEFINITIONS["YouTube"]["subscribers"],
    "source": "Sep 18 long-form snapshot plus Sep 10 reviewed Shorts catalog",
}
assert youtube_subscribers["long_form"] + youtube_subscribers["shorts"] == youtube_subscribers["gained"]
subscriber_weeks = []
subscriber_week = monday_of(min(post["date"] for post in youtube_rows))
last_subscriber_week = monday_of(max(post["date"] for post in youtube_rows))
while subscriber_week <= last_subscriber_week:
    subscriber_end = subscriber_week + timedelta(days=6)
    subscriber_coverage_end = min(subscriber_end, date.fromisoformat(YOUTUBE_UPDATE_THROUGH))
    cohort = [post for post in youtube_rows
              if subscriber_week.isoformat() <= post["date"] <= subscriber_coverage_end.isoformat()]
    subscriber_weeks.append({
        "start": subscriber_week.isoformat(),
        "end": subscriber_end.isoformat(),
        "coverage_end": subscriber_coverage_end.isoformat(),
        "covered_days": (subscriber_coverage_end - subscriber_week).days + 1,
        "gained": sum(post["followers"] for post in cohort),
        "long_form": sum(post["followers"] for post in cohort if post["type"] == "Long-Form"),
        "shorts": sum(post["followers"] for post in cohort if post["type"] == "Shorts"),
        "video_count": len(cohort),
        "complete": subscriber_end.isoformat() <= youtube_complete_through,
        "coverage_reason": (
            None if subscriber_end.isoformat() <= youtube_complete_through
            else "source-mix"
        ),
    })
    subscriber_week += timedelta(days=7)
youtube_subscribers["weekly"] = subscriber_weeks
youtube_subscribers["weekly_definition"] = (
    "Source-reported Subscribers attributed to exported videos, grouped by Monday–Sunday publish week. "
    "These are not subscribers gained during the week or the channel subscriber balance at week-end. "
    "A zero-video week means no videos in the supplied export, not zero channel activity. "
    "Weeks ending after the full Shorts cutoff are partial source-mix weeks."
)
assert sum(row["gained"] for row in subscriber_weeks) == youtube_subscribers["gained"]
assert sum(row["video_count"] for row in subscriber_weeks) == youtube_subscribers["video_count"]
assert all(row["long_form"] + row["shorts"] == row["gained"] for row in subscriber_weeks)
weekly = {}
for name, start, end in weeks:
    weekly[name] = {}
    for platform in platforms:
        subset = [x for x in posts if x["platform"] == platform and start <= x["date"] <= end]
        weekly[name][platform] = aggregate(subset)

top5 = {}
top5_coverage = {}
for name, start, end in top5_weeks:
    top5[name] = {}
    top5_coverage[name] = {}
    for platform in platforms:
        source_end = platform_ends[platform]
        complete_through = platform_complete_through[platform]
        if source_end < start:
            subset = []
            coverage_end = None
            status = "unavailable"
            covered_days = 0
            partial_reason = None
        else:
            coverage_end = min(end, source_end)
            subset = [
                x for x in posts
                if x["platform"] == platform and start <= x["date"] <= coverage_end
            ]
            status = "complete" if end <= complete_through else "partial"
            partial_reason = None if status == "complete" else (
                "source-mix" if complete_through < source_end else "date-cutoff"
            )
            covered_days = (
                date.fromisoformat(coverage_end) - date.fromisoformat(start)
            ).days + 1
        top5[name][platform] = sorted(subset, key=lambda x: (-x["views"], x["date"], x["id"]))[:5]
        top5_coverage[name][platform] = {
            "status": status,
            "source_end": source_end,
            "complete_through": complete_through,
            "coverage_end": coverage_end,
            "covered_days": covered_days,
            "partial_reason": partial_reason,
            "posts": len(subset),
        }

top5_default_week = next(
    (
        name for name, _start, _end in reversed(top5_weeks)
        if any(row["status"] == "complete" for row in top5_coverage[name].values())
        and not any(row["status"] == "partial" for row in top5_coverage[name].values())
    ),
    top5_weeks[-1][0],
)

report_name = weeks[-1][0]
report_posts = [p for p in posts if report_start.isoformat() <= p["date"] <= report_end.isoformat()]
report_totals = {p: aggregate([x for x in report_posts if x["platform"] == p]) for p in platforms}
report_scope = {
    "week": report_name,
    "start": report_start.isoformat(),
    "end": report_end.isoformat(),
    "totals": report_totals,
    "posts": len(report_posts),
    "views": sum(x["views"] for x in report_posts),
    "eng": sum(x["engagements"] for x in report_posts),
}

category_totals = {}
for platform in platforms:
    category_totals[platform] = []
    categories = sorted({p["category"] for p in posts if p["platform"] == platform})
    for category in categories:
        subset = [p for p in posts if p["platform"] == platform and p["category"] == category]
        category_totals[platform].append({"category": category, **aggregate(subset)})
    category_totals[platform].sort(key=lambda x: (-x["views"], x["category"]))

category_weekly = {}
for platform in CATEGORY_TREND_PLATFORMS:
    category_weekly[platform] = {}
    for category_row in category_totals[platform]:
        category = category_row["category"]
        category_weekly[platform][category] = {}
        for week_key, week_start, week_end in category_weeks:
            if week_start > platform_ends[platform]:
                category_weekly[platform][category][week_key] = None
                continue
            subset = [
                post for post in posts
                if post["platform"] == platform
                and post["category"] == category
                and week_start <= post["date"] <= min(week_end, platform_ends[platform])
            ]
            category_weekly[platform][category][week_key] = aggregate(subset)

tiktok_source_engagements = 0
tiktok_component_engagements = 0
tiktok_engagement_disagreements = 0
for row, _epoch in read_rows(SOURCES["TikTok"], "TT Data"):
    source_value = int(number(row.get("Total Engagements")))
    component_value = sum(
        int(number(row.get(field)))
        for field in ("Likes", "Comments", "Shares", "Bookmarks")
    )
    tiktok_source_engagements += source_value
    tiktok_component_engagements += component_value
    if source_value != component_value:
        tiktok_engagement_disagreements += 1

shorts_patch_ids = {
    clean_text(row.get("Content"))
    for row in read_csv_rows(YOUTUBE_SHORTS_LIKES_SOURCE)
    if clean_text(row.get("Content"))
}
source_notes = [
    f"YouTube: Sep 18 long-form snapshot with "
    f"{sum(p['type'] == 'Long-Form' for p in youtube_rows)} included long-form videos and "
    f"{sum(p['type'] == 'Shorts' for p in youtube_rows)} retained Shorts.",
    f"YouTube Shorts limitation: the full Shorts catalog and non-Likes metrics remain from the reviewed "
    f"Sep 10 workbook. The Sep 18 file supplies daily Likes through Sep 17 for only "
    f"{len(shorts_patch_ids)} existing Shorts; those Likes were refreshed without treating missing Shorts data as zero.",
    "YouTube subscriber metric: sum of the source-reported Subscribers column across included videos. "
    "It is not the current channel subscriber total or week-by-week channel growth.",
    "YouTube exclusions: aggregate rows, undated videos, and videos outside 2026 are omitted from dashboard analysis.",
    f"Instagram: publishing coverage through {platform_ends['Instagram']}; "
    f"{totals['Instagram']['posts']} posts, with the latest published post on {platform_latest_posts['Instagram']}. "
    "One Sep 17 Empathy Tour collaboration remains flagged in the source for authorship review and is included.",
    f"TikTok: publishing coverage through {platform_ends['TikTok']}; "
    f"{totals['TikTok']['posts']} posts, with the latest published post on {platform_latest_posts['TikTok']}. "
    f"Engagements are recomputed as Likes + Comments + Shares + Bookmarks for every row because "
    f"{tiktok_engagement_disagreements} source Total Engagements cells disagree with their components "
    f"({tiktok_source_engagements:,} source total versus {tiktok_component_engagements:,} component total).",
    f"X: combined-source publishing coverage through {platform_ends['X']}; "
    f"{totals['X']['posts']:,} authored posts, with the latest published post on {platform_latest_posts['X']}; "
    f"rows that are themselves reposted posts were excluded upstream.",
    "X engagement uses the source-reported Engagements field. Source eras can include actions beyond likes, replies, reposts, and bookmarks; Reposts interaction counts are retained.",
    "Across Instagram, TikTok, and X, recent-post metrics were refreshed Sep 18; older rows retain prior export snapshots as documented in the source Definitions sheets.",
    f"The headline reporting week uses complete cross-platform coverage through {common_coverage_end}. "
    f"Publishing coverage runs through Sep 18, but YouTube's full Shorts catalog is complete only through Sep 10.",
]

summary = {
    "brand": BRAND,
    "period": {"start": period_start, "end": period_end},
    "platform_coverage_end": platform_ends,
    "platform_complete_through": platform_complete_through,
    "platform_latest_post_date": platform_latest_posts,
    "common_coverage_end": common_coverage_end,
    "weeks": weeks,
    "top5_weeks": top5_weeks,
    "category_weeks": category_weeks,
    "report_scope": report_scope,
    "totals": totals,
    "youtube_subscribers": youtube_subscribers,
    "weekly": weekly,
    "top5": top5,
    "top5_coverage": top5_coverage,
    "top5_default_week": top5_default_week,
    "category_totals": category_totals,
    "category_weekly": category_weekly,
    "metric_definitions": METRIC_DEFINITIONS,
    "source_notes": source_notes,
    "posts": posts,
}
with open(os.path.join(ROOT, "summary.json"), "w", encoding="utf-8") as handle:
    json.dump(summary, handle, ensure_ascii=False, indent=1, sort_keys=True)

csv_path = os.path.join(ROOT, "all_posts.csv")
with open(csv_path, "w", encoding="utf-8-sig", newline="") as handle:
    writer = csv.writer(handle, lineterminator="\n")
    writer.writerow([
        "Platform", "Date", "Week", "Category", "Type", "Title / Text", "Link",
        "Views / Impressions", "Reach / Impressions", "Engagements", "Engagement Rate %",
        "Likes", "Comments / Replies", "Shares / Reposts", "Saves / Bookmarks", "Followers / Subs",
    ])
    for post in sorted(posts, key=lambda p: (-p["views"], p["platform"], p["date"])):
        rate = round(100 * post["engagements"] / post["views"], 2) if post["views"] else ""
        writer.writerow([
            post["platform"], post["date"], post["week"], post["category"], post["type"],
            post["title"], post["url"], post["views"], post["reach"], post["engagements"], rate,
            post["likes"], post["comments"], post["shares"], post["saves"], post["followers"],
        ])

print(f"=== {BRAND} · full file {period_start}..{period_end} ===")
for platform in platforms:
    total = totals[platform]
    label = METRIC_DEFINITIONS[platform]["metric_label"]
    print(f"  {platform:10s} {total['posts']:4d} posts | {total['views']:>12,} {label:11s} | {total['eng']:>9,} engagements")
print(f"--- reporting week: {report_start}..{report_end} ---")
print(f"  {report_scope['posts']} posts | {report_scope['views']:,} views/impressions | {report_scope['eng']:,} engagements")
print(f"summary.json + {csv_path}")
