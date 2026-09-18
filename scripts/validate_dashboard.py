#!/usr/bin/env python3
"""Validate generated HOTB outputs against all six current source files."""

import csv
import hashlib
import json
import os
from collections import defaultdict
from datetime import date, datetime, timedelta

from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "sources")
PATHS = {
    "youtubeBase": os.path.join(SRC, "youtube.xlsx"),
    "youtubeLong": os.path.join(SRC, "youtube_long.csv"),
    "youtubeShortsLikes": os.path.join(SRC, "youtube_shorts_likes.csv"),
    "Instagram": os.path.join(SRC, "instagram.xlsx"),
    "TikTok": os.path.join(SRC, "tiktok.xlsx"),
    "X": os.path.join(SRC, "x.xlsx"),
}


def number(value):
    if value in (None, "", "-", "--", "N/A"):
        return 0.0
    return float(value)


def clean(value):
    return str(value or "").strip()


def parse_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = clean(value)
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%Y %H:%M", "%b %d, %Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    raise ValueError(f"Unrecognized date {value!r}")


def workbook_rows(path, sheet):
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        iterator = wb[sheet].iter_rows(values_only=True)
        headers = [clean(value) for value in next(iterator)]
        return [dict(zip(headers, values)) for values in iterator if any(value is not None for value in values)]
    finally:
        wb.close()


def csv_rows(path):
    with open(path, encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def definition_date(path, label):
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        for row in wb["Definitions"].iter_rows(values_only=True):
            if clean(row[0]) == label:
                return parse_date(row[1]).isoformat()
    finally:
        wb.close()
    raise AssertionError(f"Missing {label!r} in {path}")


def sha256(path):
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def raw_total(rows, platform):
    if platform == "Instagram":
        return {
            "posts": len(rows),
            "views": int(sum(number(row["Views"]) for row in rows)),
            "eng": int(sum(sum(number(row[field]) for field in ("Likes", "Comments", "Shares", "Saves")) for row in rows)),
            "reach": int(sum(number(row["Reach"]) for row in rows)),
            "followers": int(sum(number(row["Follows"]) for row in rows)),
            "watch_hours": 0,
        }
    if platform == "TikTok":
        return {
            "posts": len(rows),
            "views": int(sum(number(row["Views"]) for row in rows)),
            "eng": int(sum(sum(number(row[field]) for field in ("Likes", "Comments", "Shares", "Bookmarks")) for row in rows)),
            "reach": 0,
            "followers": 0,
            "watch_hours": 0,
        }
    if platform == "X":
        return {
            "posts": len(rows),
            "views": int(sum(number(row["Views"]) for row in rows)),
            "eng": int(sum(number(row["Engagements"]) for row in rows)),
            "reach": 0,
            "followers": 0,
            "watch_hours": 0,
        }
    raise ValueError(platform)


with open(os.path.join(ROOT, "summary.json"), encoding="utf-8") as handle:
    summary = json.load(handle)

# Instagram, TikTok, and X raw-source contracts.
source_specs = {
    "Instagram": ("IG Data", "Post ID", "Publish time", "Permalink", "Category"),
    "TikTok": ("TT Data", "Post ID", "Publish time", "URL", "Category"),
    "X": ("X Data", "Post ID", "Date", "URL", "Category"),
}
raw = {}
for platform, (sheet, id_field, date_field, url_field, category_field) in source_specs.items():
    rows = workbook_rows(PATHS[platform], sheet)
    raw[platform] = rows
    ids = [clean(row[id_field]) for row in rows]
    urls = [clean(row[url_field]) for row in rows]
    assert ids and all(ids) and len(ids) == len(set(ids)), platform
    assert urls and all(urls) and len(urls) == len(set(urls)), platform
    assert all(row.get(date_field) and clean(row.get(category_field)) for row in rows), platform
    expected = raw_total(rows, platform)
    for key, value in expected.items():
        assert summary["totals"][platform][key] == value, (platform, key, summary["totals"][platform][key], value)

assert raw_total(raw["Instagram"], "Instagram") == {
    "posts": 659, "views": 5790528, "eng": 275801, "reach": 4467027,
    "followers": 8433, "watch_hours": 0,
}
assert raw_total(raw["TikTok"], "TikTok") == {
    "posts": 636, "views": 1686012, "eng": 81331, "reach": 0,
    "followers": 0, "watch_hours": 0,
}
assert raw_total(raw["X"], "X") == {
    "posts": 1037, "views": 114419284, "eng": 3360942, "reach": 0,
    "followers": 0, "watch_hours": 0,
}

tiktok_source_total = int(sum(number(row["Total Engagements"]) for row in raw["TikTok"]))
tiktok_component_total = raw_total(raw["TikTok"], "TikTok")["eng"]
tiktok_disagreements = [
    clean(row["Post ID"]) for row in raw["TikTok"]
    if int(number(row["Total Engagements"])) != int(sum(number(row[field]) for field in ("Likes", "Comments", "Shares", "Bookmarks")))
]
assert tiktok_source_total == 80838
assert tiktok_component_total == 81331
assert len(tiktok_disagreements) == 8

# YouTube mixed-source reconciliation.
base_shorts = [
    row for row in workbook_rows(PATHS["youtubeBase"], "Shorts Data")
    if clean(row.get("Analysis Status")) == "Included"
]
assert len(base_shorts) == 478

shorts_daily = csv_rows(PATHS["youtubeShortsLikes"])
shorts_patch = defaultdict(int)
shorts_date_keys = set()
for row in shorts_daily:
    content_id = clean(row["Content"])
    key = (content_id, parse_date(row["Date"]).isoformat())
    assert key not in shorts_date_keys
    shorts_date_keys.add(key)
    shorts_patch[content_id] += int(number(row["Likes"]))
assert len(shorts_daily) == 1300
assert len(shorts_patch) == 5
assert set(shorts_patch) <= {clean(row["Content"]) for row in base_shorts}
assert dict(shorts_patch) == {
    "-D14do6qZ3Y": 1782,
    "4QXQovg6ct8": 1643,
    "7VluUF7T5ZA": 7383,
    "GnhhMl85Q00": 1880,
    "zsK6Q-kEIUY": 4076,
}

long_all = csv_rows(PATHS["youtubeLong"])
assert len(long_all) == 322
long_detail = [row for row in long_all if clean(row["Content"]).lower() != "total"]
assert len(long_detail) == 321
assert len({clean(row["Content"]) for row in long_detail}) == 321
long_2026 = []
long_missing_date = 0
long_outside_2026 = 0
for row in long_detail:
    if not clean(row["Video publish time"]):
        long_missing_date += 1
        continue
    if parse_date(row["Video publish time"]).year != 2026:
        long_outside_2026 += 1
        continue
    long_2026.append(row)
assert (len(long_2026), long_missing_date, long_outside_2026) == (57, 245, 19)

youtube_expected = {
    "posts": len(long_2026) + len(base_shorts),
    "views": int(sum(number(row["Views"]) for row in long_2026 + base_shorts)),
    "eng": int(
        sum(number(row["Likes"]) + number(row["Shares"]) + number(row["Comments added"]) for row in long_2026)
        + sum(shorts_patch.get(clean(row["Content"]), int(number(row["Likes"])))
              + number(row["Shares"]) + number(row["Comments added"]) for row in base_shorts)
    ),
    "reach": int(sum(number(row["Thumbnail impressions"]) for row in long_2026 + base_shorts)),
    "followers": int(sum(number(row["Subscribers"]) for row in long_2026 + base_shorts)),
    "watch_hours": round(sum(number(row["Watch time (hours)"]) for row in long_2026 + base_shorts), 2),
}
assert youtube_expected == {
    "posts": 535, "views": 3406619, "eng": 141641, "reach": 2450304,
    "followers": 4382, "watch_hours": 46644.88,
}
for key, value in youtube_expected.items():
    assert summary["totals"]["YouTube"][key] == value, (key, summary["totals"]["YouTube"][key], value)

youtube_posts = [post for post in summary["posts"] if post["platform"] == "YouTube"]
assert len({post["id"] for post in youtube_posts}) == 535
assert next(post for post in youtube_posts if post["id"] == "dcCxXltUBIk")["category"] == "Bill Kristol & Tom Joscelyn"
assert next(post for post in youtube_posts if post["id"] == "j9qbcE0z_pE")["category"] == "Other"
assert next(post for post in youtube_posts if post["id"] == "tkCk8MpoYyo")["category"] == "Ad"

subscribers = summary["youtube_subscribers"]
assert subscribers["gained"] == 4382
assert subscribers["long_form"] == 1317
assert subscribers["shorts"] == 3065
assert subscribers["video_count"] == 535
assert subscribers["complete_through"] == "2026-09-10"
assert sum(row["gained"] for row in subscribers["weekly"]) == subscribers["gained"]
assert sum(row["video_count"] for row in subscribers["weekly"]) == subscribers["video_count"]

follower_metrics = summary["follower_metrics"]
assert follower_metrics["Instagram"] == {
    "available": True, "label": "Follows", "unit": "follows", "source_field": "Follows",
}
assert follower_metrics["YouTube"] == {
    "available": True, "label": "Subscribers", "unit": "subscribers", "source_field": "Subscribers",
}
assert follower_metrics["TikTok"]["available"] is False and follower_metrics["TikTok"]["source_field"] is None
assert follower_metrics["X"]["available"] is False and follower_metrics["X"]["source_field"] is None

# Freshness and calendar coverage.
expected_coverage = {"YouTube": "2026-09-18", "Instagram": "2026-09-18", "TikTok": "2026-09-18", "X": "2026-09-18"}
expected_complete = {**expected_coverage, "YouTube": "2026-09-10"}
expected_latest_post = {
    "YouTube": "2026-09-15",
    "Instagram": "2026-09-17",
    "TikTok": "2026-09-17",
    "X": "2026-09-17",
}
assert definition_date(PATHS["Instagram"], "Data through date") == "2026-09-18"
assert definition_date(PATHS["TikTok"], "Data through date") == "2026-09-18"
assert definition_date(PATHS["X"], "Data through date") == "2026-09-18"
assert definition_date(PATHS["youtubeBase"], "Data cutoff") == "2026-09-10"
assert summary["platform_coverage_end"] == expected_coverage
assert summary["platform_complete_through"] == expected_complete
assert summary["platform_latest_post_date"] == expected_latest_post
assert summary["common_coverage_end"] == "2026-09-10"
assert summary["report_scope"]["start"] == "2026-08-31"
assert summary["report_scope"]["end"] == "2026-09-06"
assert summary["report_scope"]["week"] == "Week 36"

all_posts = summary["posts"]
assert len(all_posts) == sum(total["posts"] for total in summary["totals"].values())
assert all(post["date"].startswith("2026-") and post["week"] and post["category"] for post in all_posts)
instagram_follows_by_id = {clean(row["Post ID"]): int(number(row["Follows"])) for row in raw["Instagram"]}
youtube_subscribers_by_id = {
    clean(row["Content"]): int(number(row["Subscribers"]))
    for row in long_2026 + base_shorts
}
for post in all_posts:
    if post["platform"] == "Instagram":
        assert post["followers"] == instagram_follows_by_id[post["id"]]
    elif post["platform"] == "YouTube":
        assert post["followers"] == youtube_subscribers_by_id[post["id"]]
for platform in expected_latest_post:
    assert max(post["date"] for post in all_posts if post["platform"] == platform) == expected_latest_post[platform]

report_rows = [post for post in all_posts if "2026-08-31" <= post["date"] <= "2026-09-06"]
assert summary["report_scope"]["posts"] == len(report_rows)
assert summary["report_scope"]["views"] == sum(post["views"] for post in report_rows)
assert summary["report_scope"]["eng"] == sum(post["engagements"] for post in report_rows)

weeks = summary["weeks"]
assert len(weeks) == 8
assert weeks[0] == ["Week 29", "2026-07-13", "2026-07-19"]
assert weeks[-1] == ["Week 36", "2026-08-31", "2026-09-06"]
for previous, current in zip(weeks, weeks[1:]):
    assert date.fromisoformat(current[1]) == date.fromisoformat(previous[1]) + timedelta(days=7)

category_weeks = summary["category_weeks"]
assert len(category_weeks) == 38
assert category_weeks[0] == ["2026-01-01", "2026-01-01", "2026-01-04"]
assert category_weeks[-1] == ["2026-09-14", "2026-09-14", "2026-09-18"]
for previous, current in zip(category_weeks, category_weeks[1:]):
    assert date.fromisoformat(current[1]) == date.fromisoformat(previous[2]) + timedelta(days=1)

category_keys = [row[0] for row in category_weeks]
for platform, categories in summary["category_weekly"].items():
    cutoff = summary["platform_coverage_end"][platform]
    for category, by_week in categories.items():
        assert list(by_week) == category_keys
        for key, start, end in category_weeks:
            value = by_week[key]
            if start > cutoff:
                assert value is None
                continue
            rows = [post for post in all_posts if post["platform"] == platform and post["category"] == category and start <= post["date"] <= min(end, cutoff)]
            assert value["posts"] == len(rows)
            assert value["views"] == sum(post["views"] for post in rows)
            assert value["eng"] == sum(post["engagements"] for post in rows)
            if follower_metrics[platform]["available"]:
                assert value["followers"] == sum(post["followers"] for post in rows)
        total = next(row for row in summary["category_totals"][platform] if row["category"] == category)
        metrics = ("posts", "views", "eng", "followers") if follower_metrics[platform]["available"] else ("posts", "views", "eng")
        for metric in metrics:
            assert sum(row[metric] for row in by_week.values() if row is not None) == total[metric]


def category_week_total(platform, week_key, metric):
    return sum(
        by_week[week_key][metric]
        for by_week in summary["category_weekly"][platform].values()
        if by_week[week_key] is not None
    )


assert category_week_total("Instagram", "2026-08-31", "followers") == 394
assert category_week_total("Instagram", "2026-09-07", "followers") == 3015
assert category_week_total("Instagram", "2026-09-14", "followers") == 47
assert category_week_total("YouTube", "2026-08-31", "followers") == 123
assert category_week_total("YouTube", "2026-09-07", "followers") == 81
assert category_week_total("YouTube", "2026-09-14", "followers") == 0
assert summary["category_weekly"]["Instagram"]["Produced Videos"]["2026-09-07"]["followers"] == 2656
assert summary["category_weekly"]["YouTube"]["Social Ads"]["2026-04-06"]["followers"] == -1

# Top-five coverage and ordering.
top5_weeks = summary["top5_weeks"]
assert top5_weeks[0] == ["Week 29", "2026-07-13", "2026-07-19"]
assert top5_weeks[-1] == ["Week 38", "2026-09-14", "2026-09-20"]
assert summary["top5_default_week"] == "Week 36"
for week_name, start, end in top5_weeks:
    for platform in expected_coverage:
        coverage = summary["top5_coverage"][week_name][platform]
        source_end = expected_coverage[platform]
        complete_through = expected_complete[platform]
        if source_end < start:
            expected_status, coverage_end, reason = "unavailable", None, None
        else:
            coverage_end = min(end, source_end)
            expected_status = "complete" if end <= complete_through else "partial"
            reason = None if expected_status == "complete" else ("source-mix" if complete_through < source_end else "date-cutoff")
        assert coverage["status"] == expected_status
        assert coverage["coverage_end"] == coverage_end
        assert coverage["partial_reason"] == reason
        rows = [post for post in all_posts if post["platform"] == platform and coverage_end and start <= post["date"] <= coverage_end]
        assert coverage["posts"] == len(rows)
        expected_top = sorted(rows, key=lambda post: (-post["views"], post["date"], post["id"]))[:5]
        assert [post["id"] for post in summary["top5"][week_name][platform]] == [post["id"] for post in expected_top]

w37 = summary["top5_coverage"]["Week 37"]
assert w37["YouTube"]["status"] == "partial" and w37["YouTube"]["partial_reason"] == "source-mix"
assert all(w37[platform]["status"] == "complete" for platform in ("Instagram", "TikTok", "X"))
w38 = summary["top5_coverage"]["Week 38"]
assert {platform: w38[platform]["posts"] for platform in expected_coverage} == {
    "YouTube": 1, "Instagram": 16, "TikTok": 16, "X": 13,
}
assert w38["YouTube"]["partial_reason"] == "source-mix"
assert all(w38[platform]["covered_days"] == 5 for platform in expected_coverage)

# Generated artifact checks.
dashboard_path = os.path.join(ROOT, "dashboard.html")
index_path = os.path.join(ROOT, "index.html")
dashboard = open(dashboard_path, encoding="utf-8").read()
index = open(index_path, encoding="utf-8").read()
assert dashboard == index
required_html = [
    "YouTube subscriber metric",
    "Deselect all",
    "Show all",
    "Subscribers / Follows",
    "Follower data unavailable",
    "not an account balance or a record of when audience changes occurred",
    "This is unavailable data, not a measured zero.",
    "chartDomain",
    "Posts behind this point",
    "Open video / post",
    "Partial source mix",
    "full Shorts catalog through",
    "outlined bars have partial source coverage",
    '\"top5DefaultWeek\":\"Week 36\"',
    '\"Week 38\",\"2026-09-14\",\"2026-09-20\"',
    '\"completeThrough\"',
    "Source unavailable after",
]
for text in required_html:
    assert text in dashboard, text
assert "3.4M" in dashboard or '\"views\":3406619' in dashboard
assert "114.4M" in dashboard or '\"views\":114419284' in dashboard

with open(os.path.join(ROOT, "all_posts.csv"), encoding="utf-8-sig", newline="") as handle:
    normalized = list(csv.DictReader(handle))
assert len(normalized) == len(all_posts)
for platform, expected in summary["totals"].items():
    rows = [row for row in normalized if row["Platform"] == platform]
    assert len(rows) == expected["posts"]
    assert sum(int(row["Views / Impressions"]) for row in rows) == expected["views"]
    assert sum(int(row["Engagements"]) for row in rows) == expected["eng"]
    assert sum(int(row["Followers / Subs"]) for row in rows) == expected["followers"]
assert all(row["Week"] and row["Date"] and row["Category"] for row in normalized)

result = {
    "sourceSha256": {name: sha256(path) for name, path in PATHS.items()},
    "totals": summary["totals"],
    "reportingWeek": summary["report_scope"],
    "coverage": {
        "activityThrough": summary["platform_coverage_end"],
        "completeThrough": summary["platform_complete_through"],
        "latestPost": summary["platform_latest_post_date"],
    },
    "youtube": {
        "longIncluded": len(long_2026),
        "shortsRetained": len(base_shorts),
        "shortsLikesPatched": len(shorts_patch),
        "shortsPatchMetricDates": [min(key[1] for key in shorts_date_keys), max(key[1] for key in shorts_date_keys)],
        "mixedFreshness": True,
    },
    "tiktokEngagementAudit": {
        "componentDefinitionTotal": tiktok_component_total,
        "sourceTotalEngagements": tiktok_source_total,
        "disagreementRows": tiktok_disagreements,
    },
    "topFive": {
        "first": top5_weeks[0],
        "last": top5_weeks[-1],
        "default": summary["top5_default_week"],
        "week38": w38,
    },
    "checks": [
        "All platform IDs and URLs are unique within their source",
        "Instagram, TikTok, and X totals reconcile to raw post rows",
        "TikTok engagements consistently use visible components for all rows",
        "YouTube long-form refresh reconciles to 57 dated 2026 videos",
        "Five-video Shorts feed patches only Likes and retains the 478-video catalog",
        "YouTube category mappings are preserved for matched IDs",
        "Mixed YouTube freshness is distinct from measured zero activity",
        "Headline week is the latest complete shared Monday-Sunday week",
        "Category weeks and top-five rankings reconcile to normalized posts",
        "Instagram Follows and YouTube Subscribers reconcile by post, category, and publish week",
        "TikTok and X subscriber/follow metrics are explicitly unavailable rather than zero",
        "Subscriber/follow chart supports signed YouTube adjustments",
        "Dashboard and index HTML match the normalized data layer",
        "all_posts.csv reconciles to every platform total",
    ],
}
with open(os.path.join(ROOT, "validation.json"), "w", encoding="utf-8") as handle:
    json.dump(result, handle, ensure_ascii=False, indent=2, sort_keys=True)

print(json.dumps(result, ensure_ascii=False, indent=2, sort_keys=True))
